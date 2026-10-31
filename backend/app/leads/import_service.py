"""Parse Excel workbooks and create leads for a project."""

from __future__ import annotations

import io
import logging
import re
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.common.exceptions import AppError, ErrorCode
from app.lead_fields.models import LeadField
from app.lead_fields.service import active_fields_for_project
from app.leads.models import Lead
from app.leads.schemas import (
    CreateLeadRequest,
    ImportLeadError,
    ImportLeadsResult,
    LeadPriority,
    normalize_priority,
    normalize_status,
)
from app.leads.service import create_lead
from app.organization.models import Project
from app.pipeline.service import assert_stage_key, entry_stage_key
from app.users.models import User

logger = logging.getLogger("app.leads.import")

_MAX_BYTES = 5 * 1024 * 1024
_MAX_ROWS = 500

_COMMON_ALIASES: dict[str, str] = {
    "name": "full_name",
    "contact person": "full_name",
    "full name": "full_name",
    "phone": "mobile",
    "mobile": "mobile",
    "mobile number": "mobile",
    "mobile no": "mobile",
    "mobile no.": "mobile",
    "contact no": "mobile",
    "contact no.": "mobile",
    "email": "email",
    "email id": "email",
    "city": "city",
    "location": "city",
    "company": "company_name",
    "company name": "company_name",
    "school name": "company_name",
    "requirement": "requirement",
    "requirement type": "requirement",
    "requirement details": "requirement",
    "lead source": "lead_source",
    "source": "lead_source",
    "amount": "expected_value",
    "expected value": "expected_value",
    "order amount": "expected_value",
    "priority": "priority",
    "warm": "priority",
    "prority": "priority",
    "lead status": "status",
    "status": "status",
    "converted": "status",
    "next follow-up": "next_followup_at",
    "next follow up": "next_followup_at",
    "follow-up date": "next_followup_at",
    "follow up date": "next_followup_at",
}

_NOTE_ALIASES = frozenset(
    {
        "call summary",
        "follow-up",
        "follow up",
        "follow-up notes",
        "follow up notes",
        "followup comments",
        "followup comments ",
        "discussion summary",
        "remarks",
    }
)

_STATUS_ALIASES = {
    "followup": "in_progress",
    "follow-up": "in_progress",
    "follow up": "in_progress",
    "converted": "won",
    "closed": "won",
    "connected": "contacted",
    "not connected": "contacted",
    "lost": "lost",
    "new": "new",
    "hot": "new",
    "warm": "contacted",
    "cold": "nurture",
}

_PRIORITY_ALIASES = {
    "hot": "high",
    "warm": "medium",
    "cold": "low",
    "high": "high",
    "medium": "medium",
    "low": "low",
}


def _norm_header(value: object) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).strip().lower().split())
    return text.rstrip(".")


def _is_note_header(header: str) -> bool:
    if not header:
        return False
    if header.startswith("next follow"):
        return False
    if header in _NOTE_ALIASES:
        return True
    if header.startswith("call summary"):
        return True
    if header.startswith("followup comment") or header.startswith("follow-up comment"):
        return True
    if header.startswith("discussion summary"):
        return True
    return False


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _cell_phone(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        text = _cell_text(value)
        return None if text is None else "".join(ch for ch in text if ch.isdigit()) or None
    if isinstance(value, int):
        return str(value)
    text = _cell_text(value)
    if text is None:
        return None
    digits = "".join(ch for ch in text if ch.isdigit())
    return digits or None


def _parse_decimal(value: object) -> Decimal | None:
    text = _cell_text(value)
    if text is None:
        return None
    cleaned = text.replace(",", "").replace("₹", "").replace("Rs", "").replace("rs", "").strip()
    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    if amount < 0:
        return None
    return amount


def _parse_datetime(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, 12, 0, tzinfo=UTC)
    text = _cell_text(value)
    if text is None:
        return None
    for fmt in (
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y %I:%M %p",
        "%d/%m/%Y - %I:%M %p",
        "%d/%m/%Y - %H:%M",
    ):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})", text)
    if match:
        day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
        if year < 100:
            year += 2000
        try:
            return datetime(year, month, day, 12, 0, tzinfo=UTC)
        except ValueError:
            return None
    return None


def _map_priority(raw: object) -> LeadPriority:
    text = (_cell_text(raw) or "").strip().lower()
    mapped = _PRIORITY_ALIASES.get(text)
    if mapped in ("low", "medium", "high"):
        return mapped  # type: ignore[return-value]
    return normalize_priority(text or "medium")  # type: ignore[return-value]


def _map_status_raw(raw: object) -> str | None:
    text = (_cell_text(raw) or "").strip().lower()
    if not text:
        return None
    if text in _STATUS_ALIASES:
        return _STATUS_ALIASES[text]
    return normalize_status(text)


def _notes_field_key(fields: list[LeadField]) -> str | None:
    by_label = {_norm_header(field.label): field.key for field in fields}
    for label in ("follow-up notes", "follow up notes", "call summary", "remarks"):
        if label in by_label:
            return by_label[label]
    for key in ("follow_up_notes", "followup_notes", "call_summary", "remarks"):
        if any(field.key == key for field in fields):
            return key
    return None


def _extra_key_for_header(header: str, fields: list[LeadField]) -> str | None:
    if not header or header in _COMMON_ALIASES or _is_note_header(header):
        return None
    # Skip identifiers / unused
    if header in {"lead id", "sr.no", "sr no", "date"}:
        return None
    by_label = {_norm_header(field.label): field.key for field in fields}
    if header in by_label:
        return by_label[header]
    # soft aliases for Paneerwala
    soft = {
        "qty": "quantity",
        "quantity": "quantity",
        "product": "product",
        "payment": "payment",
        "call time": "call time",
        "enquiry received": "enquiry date",
        "enquiry date": "enquiry date",
        "order status": "order status",
        "ordering information": "ordering information",
        "interested product": "interested product",
    }
    target_label = soft.get(header, header)
    if target_label in by_label:
        return by_label[target_label]
    # match by key slug-ish
    slug = header.replace(" ", "_").replace("-", "_")
    for field in fields:
        if field.key == slug:
            return field.key
    return None


def _default_followup() -> datetime:
    tomorrow = datetime.now(UTC).date() + timedelta(days=1)
    return datetime(tomorrow.year, tomorrow.month, tomorrow.day, 12, 0, tzinfo=UTC)


def _find_existing_mobile(db: Session, company_id: uuid.UUID, mobile: str) -> bool:
    digits = "".join(ch for ch in mobile if ch.isdigit())
    if not digits:
        return False
    rows = db.scalars(
        select(Lead.mobile).where(
            Lead.company_id == company_id,
            Lead.is_deleted.is_(False),
        )
    ).all()
    for existing in rows:
        existing_digits = "".join(ch for ch in (existing or "") if ch.isdigit())
        if existing_digits and existing_digits == digits:
            return True
    return False


def import_leads_from_excel(
    db: Session,
    *,
    actor: User,
    project_id: uuid.UUID,
    file_bytes: bytes,
    filename: str,
    ip_address: str | None,
) -> ImportLeadsResult:
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise AppError(
            ErrorCode.VALIDATION_ERROR,
            "Upload an Excel .xlsx file.",
            422,
            details=[{"field": "file", "code": "INVALID_TYPE", "message": "Upload an Excel .xlsx file."}],
        )
    if len(file_bytes) > _MAX_BYTES:
        raise AppError(
            ErrorCode.VALIDATION_ERROR,
            "File is too large (max 5 MB).",
            422,
            details=[{"field": "file", "code": "TOO_LARGE", "message": "File is too large (max 5 MB)."}],
        )
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise AppError(
            ErrorCode.VALIDATION_ERROR,
            "Could not read that Excel file.",
            422,
            details=[{"field": "file", "code": "INVALID_FILE", "message": "Could not read that Excel file."}],
        ) from exc

    sheet = workbook[workbook.sheetnames[0]]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise AppError(
            ErrorCode.VALIDATION_ERROR,
            "The spreadsheet is empty.",
            422,
            details=[{"field": "file", "code": "EMPTY", "message": "The spreadsheet is empty."}],
        ) from None

    headers = [_norm_header(cell) for cell in header_row]
    if not any(headers):
        raise AppError(
            ErrorCode.VALIDATION_ERROR,
            "The first row must contain column headers.",
            422,
            details=[{"field": "file", "code": "NO_HEADERS", "message": "The first row must contain column headers."}],
        )

    fields = active_fields_for_project(db, project_id)
    notes_key = _notes_field_key(fields)
    column_common: dict[int, str] = {}
    column_extra: dict[int, str] = {}
    note_indexes: list[int] = []
    for index, header in enumerate(headers):
        if not header:
            continue
        if _is_note_header(header):
            note_indexes.append(index)
            continue
        common = _COMMON_ALIASES.get(header)
        if common:
            column_common[index] = common
            continue
        extra_key = _extra_key_for_header(header, fields)
        if extra_key:
            column_extra[index] = extra_key

    created = 0
    skipped = 0
    errors: list[ImportLeadError] = []
    company_id: uuid.UUID | None = None

    for row_offset, row in enumerate(rows_iter, start=2):
        if created + skipped + len(errors) >= _MAX_ROWS:
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message=f"Stopped after {_MAX_ROWS} rows.",
                )
            )
            break
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        values: dict[str, object] = {}
        for index, field_name in column_common.items():
            if index < len(row) and row[index] is not None and str(row[index]).strip() != "":
                values[field_name] = row[index]

        full_name = _cell_text(values.get("full_name"))
        mobile = _cell_phone(values.get("mobile"))
        if not full_name and not mobile:
            continue
        if not full_name or not mobile:
            skipped += 1
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message="Name and phone are required.",
                    name=full_name,
                    mobile=mobile,
                )
            )
            continue

        fields_by_key = {field.key: field for field in fields}
        extras: dict[str, Any] = {}
        for index, key in column_extra.items():
            if index >= len(row):
                continue
            text = _cell_text(row[index])
            if text is None:
                continue
            field_def = fields_by_key.get(key)
            if field_def is not None and field_def.field_type == "number":
                match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
                if match:
                    extras[key] = match.group(0)
                continue
            if field_def is not None and field_def.field_type == "date":
                parsed = _parse_datetime(row[index])
                if parsed is not None:
                    extras[key] = parsed.date().isoformat()
                continue
            extras[key] = text

        note_parts: list[str] = []
        for index in note_indexes:
            if index >= len(row):
                continue
            text = _cell_text(row[index])
            if text:
                note_parts.append(text)
        if note_parts:
            merged = "\n\n".join(note_parts)
            if notes_key:
                extras[notes_key] = merged
            else:
                existing_req = _cell_text(values.get("requirement")) or ""
                values["requirement"] = f"{existing_req}\n\n{merged}".strip() if existing_req else merged

        priority = _map_priority(values.get("priority"))
        status_raw = _map_status_raw(values.get("status"))
        follow_at = _parse_datetime(values.get("next_followup_at")) or _default_followup()
        expected = _parse_decimal(values.get("expected_value"))
        email = _cell_text(values.get("email"))
        city = _cell_text(values.get("city"))
        company_name = _cell_text(values.get("company_name"))
        requirement = _cell_text(values.get("requirement"))
        lead_source = _cell_text(values.get("lead_source"))

        # Resolve status against pipeline; fall back to entry stage.
        status_for_create: str | None = None
        if status_raw:
            try:
                status_for_create = assert_stage_key(db, project_id, status_raw)
            except AppError:
                status_for_create = entry_stage_key(db, project_id)

        try:
            payload = CreateLeadRequest(
                full_name=full_name,
                mobile=mobile,
                email=email,
                company_name=company_name,
                city=city,
                requirement=requirement,
                expected_value=expected,
                lead_source=lead_source,
                project_id=project_id,
                status=status_for_create,
                priority=priority,
                next_followup_at=follow_at,
                extra_values=extras or None,
            )
        except ValidationError as exc:
            skipped += 1
            first = exc.errors()[0] if exc.errors() else {}
            message = str(first.get("msg") or "Invalid row data.")
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message=message,
                    name=full_name,
                    mobile=mobile,
                )
            )
            continue

        project = db.get(Project, project_id)
        if project is None:
            raise AppError(ErrorCode.VALIDATION_ERROR, "Choose a project.", 422)
        company_id = project.company_id
        if _find_existing_mobile(db, company_id, mobile):
            skipped += 1
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message="Skipped — mobile already exists.",
                    name=full_name,
                    mobile=mobile,
                )
            )
            continue

        try:
            create_lead(db, actor=actor, payload=payload, ip_address=ip_address)
            created += 1
        except AppError as exc:
            skipped += 1
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message=exc.message,
                    name=full_name,
                    mobile=mobile,
                )
            )
        except Exception as exc:  # noqa: BLE001
            logger.exception("import_row_failed", extra={"row": row_offset})
            skipped += 1
            errors.append(
                ImportLeadError(
                    row=row_offset,
                    message="Could not import this row.",
                    name=full_name,
                    mobile=mobile,
                )
            )

    workbook.close()
    return ImportLeadsResult(
        created=created,
        skipped=skipped,
        errors=errors[:50],
        headers_seen=[h for h in headers if h],
    )
